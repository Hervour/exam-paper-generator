from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(r"C:\Users\LHW66\Documents\Codex\2026-09-16\bang")
OUT = BASE / "outputs"
OUT.mkdir(parents=True, exist_ok=True)

BODY_FONT_EA = "宋体"
HEAD_FONT_EA = "微软雅黑"
BODY_FONT_LATIN = "Times New Roman"
HEAD_FONT_LATIN = "Microsoft YaHei"
DARK_BLUE = "1F4E79"
MID_BLUE = "2E74B5"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"


def set_run_font(run, size=10.5, bold=None, color=None, italic=None,
                 east_asia=BODY_FONT_EA, latin=BODY_FONT_LATIN):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), latin)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, size, bold=False, color=None,
                   east_asia=BODY_FONT_EA, latin=BODY_FONT_LATIN):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), latin)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = f"w:{edge}"
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key in ["sz", "val", "color", "space"]:
                if key in edge_data:
                    element.set(qn(f"w:{key}"), str(edge_data[key]))


def add_runs(paragraph, text, size=10.5, color=None, base_bold=False):
    parts = text.split("**")
    for idx, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        set_run_font(run, size=size, bold=(base_bold or idx % 2 == 1), color=color)


def add_body(doc, text, space_after=4):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(space_after)
    add_runs(p, text, size=10.5)
    return p


def add_bullet(doc, text, level=0, space_after=3):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Cm(0.74 + 0.65 * level)
    p.paragraph_format.first_line_indent = Cm(-0.42)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.2
    add_runs(p, text, size=10.5)
    return p


def add_formula(doc, text, space_after=5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_font(run, size=11, east_asia=BODY_FONT_EA, latin="Cambria Math")
    return p


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(9)
    run = p.add_run(text)
    set_run_font(run, size=18, bold=True, color=DARK_BLUE,
                 east_asia=HEAD_FONT_EA, latin=HEAD_FONT_LATIN)
    return p


def add_heading(doc, text, level=1):
    style = "Heading 1" if level == 1 else "Heading 2"
    p = doc.add_paragraph(style=style)
    if level == 1:
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(5)
        size = 13.5
    else:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        size = 11.5
    p.paragraph_format.keep_with_next = True
    add_runs(p, text, size=size, color=(DARK_BLUE if level == 1 else MID_BLUE), base_bold=True)
    return p


def add_table(doc, headers, rows, widths_cm, header_fill=LIGHT_BLUE, first_col_left=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for idx, (cell, text) in enumerate(zip(hdr.cells, headers)):
        cell.width = Cm(widths_cm[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        add_runs(p, text, size=9.5, base_bold=True, color=DARK_BLUE)
        set_cell_border(cell,
            top={"val": "single", "sz": "8", "color": "7F9DB9"},
            bottom={"val": "single", "sz": "8", "color": "7F9DB9"},
            left={"val": "single", "sz": "6", "color": "9CB3C9"},
            right={"val": "single", "sz": "6", "color": "9CB3C9"})

    for row_idx, row_values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for col_idx, (cell, text) in enumerate(zip(row.cells, row_values)):
            cell.width = Cm(widths_cm[col_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if first_col_left and col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, str(text), size=9.3)
            set_cell_border(cell,
                top={"val": "single", "sz": "6", "color": "A9BCD0"},
                bottom={"val": "single", "sz": "6", "color": "A9BCD0"},
                left={"val": "single", "sz": "6", "color": "A9BCD0"},
                right={"val": "single", "sz": "6", "color": "A9BCD0"})

    for col_idx, width in enumerate(widths_cm):
        for cell in table.columns[col_idx].cells:
            cell.width = Cm(width)
    return table


def configure_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.85)
    sec.bottom_margin = Cm(1.75)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    sec.header_distance = Cm(1.0)
    sec.footer_distance = Cm(1.0)

    normal = doc.styles["Normal"]
    set_style_font(normal, 10.5, east_asia=BODY_FONT_EA, latin=BODY_FONT_LATIN)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(4)

    for name, size, color in (("Heading 1", 13.5, DARK_BLUE), ("Heading 2", 11.5, MID_BLUE)):
        style = doc.styles[name]
        set_style_font(style, size, bold=True, color=color,
                       east_asia=HEAD_FONT_EA, latin=HEAD_FONT_LATIN)

    for name in ("List Bullet", "List Bullet 2"):
        style = doc.styles[name]
        set_style_font(style, 10.5, east_asia=BODY_FONT_EA, latin=BODY_FONT_LATIN)


def build_first():
    doc = Document()
    configure_document(doc)
    add_title(doc, "第一题：生产线月度成本构成分析")
    add_body(doc, "按每月 30 天计算，月产量为：", space_after=2)
    add_formula(doc, "1,000 × 30 = 30,000（台/月）", space_after=6)

    add_heading(doc, "基础成本数据", 1)
    add_bullet(doc, "**生产线折旧费**：2 条生产线总值 60 万元，按 10 年（120 个月）直线折旧，月折旧为：600,000 ÷ 120 = 5,000 元。")
    add_bullet(doc, "**生产线人工**：2 × 3 × 5,000 = 30,000 元/月。")
    add_bullet(doc, "**厂房水电煤**：200,000 元/月。")
    add_bullet(doc, "**库房管理人员**：2 × 3,000 = 6,000 元/月。")
    add_bullet(doc, "**单台变动材料与装配**：每台 500 + 100 = 600 元；月总计 600 × 30,000 = 18,000,000 元。")

    add_heading(doc, "1. 财务会计视角（完全成本法）", 1)
    add_body(doc, "按经济用途分为制造成本与非制造成本。")
    add_heading(doc, "制造成本（计入产品成本）", 2)
    add_bullet(doc, "直接材料：500 元/台 × 30,000 = 15,000,000 元。")
    add_bullet(doc, "直接人工/装配费：装配费用 100 元/台 × 30,000 = 3,000,000 元。")
    add_bullet(doc, "制造费用：生产线折旧（5,000 元）+ 产线维护人工（30,000 元）+ 厂房水电能源费（200,000 元）= 235,000 元。")
    add_heading(doc, "非制造成本（期间费用）", 2)
    add_bullet(doc, "管理费用/仓储费用：库房管理员人工成本 6,000 元。")

    add_heading(doc, "2. 管理会计视角（成本性态/变动成本法）", 1)
    add_body(doc, "按成本与产量的依存关系分类。")
    add_heading(doc, "变动成本", 2)
    add_bullet(doc, "总额随产量正比例变动，包括材料费 15,000,000 元与装配费 3,000,000 元，合计 18,000,000 元（单位变动成本 600 元/台）。")
    add_heading(doc, "固定成本", 2)
    add_bullet(doc, "在产能范围内不受产量影响，包括生产线折旧 5,000 元、产线固定维护人工 30,000 元、厂房基础能耗及库房管理人工 6,000 元。")
    add_heading(doc, "混合成本", 2)
    add_bullet(doc, "水电煤能源费 200,000 元通常包含照明等维持基本运营的固定能耗与随产量增减的机器动力能耗；管理会计中需依据作业动因进一步拆分为固定与变动两部分。")

    path = OUT / "第一题.docx"
    doc.save(path)
    return path


def build_second():
    doc = Document()
    configure_document(doc)
    add_title(doc, "第二题：作业成本法计算与对比分析")

    add_heading(doc, "1. 计算各作业动因分配率", 1)
    add_table(
        doc,
        ["作业活动", "制造费用总成本（元）", "成本动因", "总动因量", "动因分配率"],
        [
            ["机器调试", "600,000", "调试次数", "30 + 120 = 150 次", "600,000 ÷ 150 = **4,000 元/次**"],
            ["质量检验", "500,000", "检验批次", "40 + 110 = 150 批次", "500,000 ÷ 150 = **3,333.33 元/批次**"],
            ["订单处理", "400,000", "订单数量", "20 + 60 = 80 单", "400,000 ÷ 80 = **5,000 元/单**"],
            ["设备维护", "1,300,000", "机器工时", "30,000 + 20,000 = 50,000 小时", "1,300,000 ÷ 50,000 = **26 元/小时**"],
        ],
        [2.3, 3.4, 2.4, 3.7, 5.2],
    )

    add_heading(doc, "2. 制造费用分配计算", 1)
    add_heading(doc, "标准型电路板", 2)
    add_bullet(doc, "机器调试：30 × 4,000 = 120,000 元")
    add_bullet(doc, "质量检验：40 × 3,333.33 = 133,333.33 元")
    add_bullet(doc, "订单处理：20 × 5,000 = 100,000 元")
    add_bullet(doc, "设备维护：30,000 × 26 = 780,000 元")
    p = add_body(doc, "制造费用合计：**1,133,333.33 元**", space_after=5)
    p.paragraph_format.left_indent = Cm(0.74)

    add_heading(doc, "定制型电路板", 2)
    add_bullet(doc, "机器调试：120 × 4,000 = 480,000 元")
    add_bullet(doc, "质量检验：110 × 3,333.33 = 366,666.67 元")
    add_bullet(doc, "订单处理：60 × 5,000 = 300,000 元")
    add_bullet(doc, "设备维护：20,000 × 26 = 520,000 元")
    p = add_body(doc, "制造费用合计：**1,666,666.67 元**", space_after=5)
    p.paragraph_format.left_indent = Cm(0.74)

    add_heading(doc, "3. 产品总成本与单位成本", 1)
    add_table(
        doc,
        ["成本项目", "标准型电路板（年产 10,000 块）", "定制型电路板（年产 2,000 块）"],
        [
            ["直接材料", "2,000,000.00 元", "800,000.00 元"],
            ["直接人工", "600,000.00 元", "200,000.00 元"],
            ["分配的制造费用", "1,133,333.33 元", "1,666,666.67 元"],
            ["**总成本**", "**3,733,333.33 元**", "**2,666,666.67 元**"],
            ["**单位成本**", "**373.33 元/块**", "**1,333.33 元/块**"],
        ],
        [4.0, 6.4, 6.4],
    )

    add_heading(doc, "4. 成本对比分析", 1)
    add_heading(doc, "成本倒挂与资源消耗动因", 2)
    add_body(doc, "定制型电路板产量仅占总量的 16.7%，却分摊了约 59.5%（1,666,666.67 元）的制造费用。原因在于定制型产品具有小批量、高复杂度的特点，过度消耗了机器调试（占比 80%）、质量检验（占比 73.3%）和订单处理（占比 75%）等非产量驱动作业。")
    add_heading(doc, "传统成本法 vs 作业成本法（ABC）", 2)
    add_body(doc, "若按传统方法以直接人工工时或产量分配制造费用，大批量的标准型产品会替定制型产品承担过多间接费用，导致“标准型高估成本、定制型低估成本”的成本扭曲现象。采用作业成本法能真实反映产品消耗资源的因果关系，为定制产品的合理定价与产线优化决策提供准确依据。")

    path = OUT / "第二题.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_first())
    print(build_second())
