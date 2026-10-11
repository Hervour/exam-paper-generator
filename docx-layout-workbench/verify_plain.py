from pathlib import Path
from docx import Document
import pdfplumber, re

base=Path(r'C:\Users\LHW66\Documents\Codex\2026-09-16\bang')
for name in ('第一题','第二题'):
    doc=Document(base/'outputs'/f'{name}.docx')
    runs=[]
    for p in doc.paragraphs:
        runs.extend(p.runs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    runs.extend(p.runs)
    nonblack=[]
    for r in runs:
        rgb=r.font.color.rgb
        if rgb is not None and str(rgb).upper() not in ('000000','FFFFFF'):
            nonblack.append((r.text[:20], str(rgb)))
    numbered=[]
    def check_para(p):
        t=p.text.strip()
        if re.match(r'^\d+\.\s', t):
            numbered.append(t[:40])
    for p in doc.paragraphs:
        check_para(p)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    check_para(p)
    assert not nonblack, nonblack[:10]
    assert not numbered, numbered[:10]
    pdf_path=base/'work'/'qa_plain'/name/f'{name}.pdf'
    with pdfplumber.open(pdf_path) as pdf:
        text_colors=[]; line_colors=[]; overlaps=[]; page_bottoms=[]
        for page in pdf.pages:
            chars=page.chars
            page_bottoms.append(max((c['bottom'] for c in chars), default=0))
            for c in chars:
                col=c.get('non_stroking_color')
                if col is not None: text_colors.append(tuple(col) if isinstance(col,(list,tuple)) else col)
            for line in page.lines:
                col=line.get('stroking_color')
                if col is not None: line_colors.append(tuple(col) if isinstance(col,(list,tuple)) else col)
            words=page.extract_words()
            for a in range(len(words)):
                wa=words[a]
                for b in range(a+1,len(words)):
                    wb=words[b]
                    x0=max(wa['x0'],wb['x0']); x1=min(wa['x1'],wb['x1'])
                    y0=max(wa['top'],wb['top']); y1=min(wa['bottom'],wb['bottom'])
                    if x1>x0 and y1>y0:
                        area=(x1-x0)*(y1-y0)
                        min_area=min((wa['x1']-wa['x0'])*(wa['bottom']-wa['top']),(wb['x1']-wb['x0'])*(wb['bottom']-wb['top']))
                        if min_area and area/min_area>0.35: overlaps.append((wa['text'],wb['text']))
        unique_text=sorted(set(text_colors), key=str)
        unique_line=sorted(set(line_colors), key=str)
        print(f'{name}: runs={len(runs)}, pages={len(pdf.pages)}, text_colors={unique_text}, line_colors={unique_line[:12]}, overlaps={overlaps[:3]}, last_bottom={max(page_bottoms):.1f}')
