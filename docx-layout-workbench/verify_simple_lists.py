from pathlib import Path
from docx import Document
base=Path(r'C:\Users\LHW66\Documents\Codex\2026-09-16\bang\outputs')/'修改版'
for name in ('第一题','第二题'):
    path=base/f'{name}.docx'
    doc=Document(path)
    texts=[p.text.strip() for p in doc.paragraphs if p.text.strip()]
    bullet=[t for t in texts if t.startswith(('•','-','*')) or '•' in t]
    numbered=[t for t in texts if len(t)>3 and t[0].isdigit() and t[1:3] in ('. ','．')]
    assert not bullet, bullet
    print(f'{name}: no bullet characters; numbered paragraphs={len(numbered)}; tables={len(doc.tables)}')
