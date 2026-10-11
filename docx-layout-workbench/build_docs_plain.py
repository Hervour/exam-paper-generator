from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(r"C:\Users\LHW66\Documents\Codex\2026-09-16\bang")
OUT = BASE / "outputs"
OUT.mkdir(parents=True, exist_ok=True)

EA_BODY = "宋体"
EA_HEAD = "黑体"
LATIN = "Times New Roman"


def set_run_font(run, size=10.5, bold=None, color=RGBColor(0, 0, 0),
                 east_asia=EA_BODY, latin=LATIN):
    run.font.name = latin
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), latin)


def set_style_font(style, size, bold=False, color=RGBColor(0, 0, 0),
                   east_asia=EA_BODY, latin=LATIN):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), latin)


def add_runs(paragraph, text, size=10.5, base_bold=False):
    parts = text.split("**")
    for idx, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        set_run_font(run, size=size, bold=(base_bold or idx % 2 == 1))


def add_body(doc, text, space_after=3):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(space_after)
    add_runs(p, text)
    return p


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=16, bold=True, east_asia=EA_HEAD)


def add_heading(doc, text, level=1):
    style = "Heading 1" if level == 1 else "Heading 2"
    p = doc.add_paragraph(style=style)
    if level == 1:
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
        size = 12.5
    else:
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        size = 11
    p.paragraph_format.keep_with_next = True
    add_runs(p, text, size=size, base_bold=True)
    return p


def add_formula(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    return p


def add_number(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.first_line_indent = Cm(-0.55)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.25
    add_runs(p, text)
    return p


def add_numbered(doc, items):
    for number, text in enumerate(items, 1):
        add_number(doc, number, text)


def set_cell_margins(cell, top=60, start=80, bottom=60, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_table(doc, headers, rows, widths_cm):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    prevent_row_split(header_row)
    for idx, cell in enumerate(header_row.cells):
        cell.width = Cm(widths_cm[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        add_runs(p, headers[idx], size=9.5, base_bold=True)

    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for idx, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_runs(p, row_values[idx], size=9.3)

    for col_idx, width in enumerate(widths_cm):
        for cell in table.columns[col_idx].cells:
            cell.width = Cm(width)
    return table


def configure_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    set_style_font(normal, 10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(3)

    set_style_font(doc.styles["Heading 1"], 12.5, bold=True, east_asia=EA_HEAD)
    set_style_font(doc.styles["Heading 2"], 11, bold=True, east_asia=EA_HEAD)


def build_first():
    doc = Document()
    configure_document(doc)
    add_title(doc, "第一题：生产线月度成本构成分析")
    add_body(doc, "按每月 30 天计算，月产量为：")
    add_formula(doc, "1,000 × 30 = 30,000（台/月）")

    add_heading(doc, "基础成本数据", 1)
    add_numbered(doc, [
        "**生产线折旧费**：2 条生产线总值 60 万元，按 10 年（120 个月）直线折旧，月折旧为 600,000 ÷ 120 = 5,000 元。",
        "**生产线人工**：2 × 3 × 5,000 = 30,000 元/月。",
        "**厂房水电煤**：200,000 元/月。",
        "**库房管理人员**：2 × 3,000 = 6,000 元/月。",
        "**单台变动材料与装配**：每台 500 + 100 = 600 元，月总计 600 × 30,000 = 18,000,000 元。",
    ])

    add_heading(doc, "财务会计视角（完全成本法）", 1)
    add_body(doc, "按经济用途分为制造成本与非制造成本。")
    add_heading(doc, "制造成本（计入产品成本）", 2)
    add_numbered(doc, [
        "直接材料：500 元/台 × 30,000 = 15,000,000 元。",
        "直接人工/装配费：装配费用 100 元/台 × 30,000 = 3,000,000 元。",
        "制造费用：生产线折旧 5,000 元 + 产线维护人工 30,000 元 + 厂房水电能源费 200,000 元 = 235,000 元。",
    ])
    add_heading(doc, "非制造成本（期间费用）", 2)
    add_number(doc, 1, "管理费用/仓储费用：库房管理员人工成本 6,000 元。")

    add_heading(doc, "管理会计视角（成本性态/变动成本法）", 1)
    add_body(doc, "按成本与产量的依存关系分类。")
    add_heading(doc, "变动成本", 2)
    add_number(doc, 1, "总额随产量正比例变动，包括材料费 15,000,000 元与装配费 3,000,000 元，合计 18,000,000 元，单位变动成本 600 元/台。")
    add_heading(doc, "固定成本", 2)
    add_number(doc, 1, "在产能范围内不受产量影响，包括生产线折旧 5,000 元、产线固定维护人工 30,000 元、厂房基础能耗及库房管理人工 6,000 元。")
    add_heading(doc, "混合成本", 2)
    add_number(doc, 1, "水电煤能源费 200,000 元通常包含照明等维持基本运营的固定能耗与随产量增减的机器动力能耗，管理会计中需依据作业动因进一步拆分为固定与变动两部分。")

    path = OUT / "第一题.docx"
    doc.save(path)
    return path


def build_second():
    doc = Document()
    configure_document(doc)
    add_title(doc, "第二题：作业成本法计算与对比分析")

    add_heading(doc, "计算各作业动因分配率", 1)
    add_table(
        doc,
        ["作业活动", "制造费用总成本（元）", "成本动因", "总动因量", "动因分配率"],
        [
            ["机器调试", "600,000", "调试次数", "30 + 120 = 150 次", "600,000 ÷ 150 = 4,000 元/次"],
            ["质量检验", "500,000", "检验批次", "40 + 110 = 150 批次", "500,000 ÷ 150 = 3,333.33 元/批次"],
            ["订单处理", "400,000", "订单数量", "20 + 60 = 80 单", "400,000 ÷ 80 = 5,000 元/单"],
            ["设备维护", "1,300,000", "机器工时", "30,000 + 20,000 = 50,000 小时", "1,300,000 ÷ 50,000 = 26 元/小时"],
        ],
        [2.3, 3.4, 2.4, 3.7, 5.2],
    )

    add_heading(doc, "制造费用分配计算", 1)
    add_heading(doc, "标准型电路板", 2)
    add_numbered(doc, [
        "机器调试：30 × 4,000 = 120,000 元。",
        "质量检验：40 × 3,333.33 = 133,333.33 元。",
        "订单处理：20 × 5,000 = 100,000 元。",
        "设备维护：30,000 × 26 = 780,000 元。",
        "制造费用合计：1,133,333.33 元。",
    ])
    add_heading(doc, "定制型电路板", 2)
    add_numbered(doc, [
        "机器调试：120 × 4,000 = 480,000 元。",
        "质量检验：110 × 3,333.33 = 366,666.67 元。",
        "订单处理：60 × 5,000 = 300,000 元。",
        "设备维护：20,000 × 26 = 520,000 元。",
        "制造费用合计：1,666,666.67 元。",
    ])

    add_heading(doc, "产品总成本与单位成本", 1)
    add_table(
        doc,
        ["成本项目", "标准型电路板（年产 10,000 块）", "定制型电路板（年产 2,000 块）"],
        [
            ["直接材料", "2,000,000.00 元", "800,000.00 元"],
            ["直接人工", "600,000.00 元", "200,000.00 元"],
            ["分配的制造费用", "1,133,333.33 元", "1,666,666.67 元"],
            ["总成本", "3,733,333.33 元", "2,666,666.67 元"],
            ["单位成本", "373.33 元/块", "1,333.33 元/块"],
        ],
        [4.0, 6.3, 6.3],
    )

    add_heading(doc, "成本对比分析", 1)
    add_heading(doc, "成本倒挂与资源消耗动因", 2)
    add_body(doc, "定制型电路板产量仅占总量的 16.7%，却分摊了约 59.5%（1,666,666.67 元）的制造费用。原因在于定制型产品具有小批量、高复杂度的特点，过度消耗了机器调试（占比 80%）、质量检验（占比 73.3%）和订单处理（占比 75%）等非产量驱动作业。")
    add_heading(doc, "传统成本法 vs 作业成本法（ABC）", 2)
    add_body(doc, "若按传统方法以直接人工工时或产量分配制造费用，大批量的标准型产品会替定制型产品承担过多间接费用，导致“标准型高估成本、定制型低估成本”的成本扭曲现象。采用作业成本法能真实反映产品消耗资源的因果关系，为定制产品的合理定价与产线优化决策提供准确依据。")

    path = OUT / "第二题.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_first())
    print(build_second())

