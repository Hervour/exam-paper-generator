# -*- coding: utf-8 -*-
import docx, os
vault = r"C:\obsidian知识库\暑期课程\分班考试模拟卷"
desktop = r"C:\Users\LHW66\Desktop"
print("vault files:", sorted(os.listdir(vault)))
for name in sorted(os.listdir(desktop)):
    if name.endswith(".docx") and "分班考试模拟卷" in name:
        d = docx.Document(os.path.join(desktop, name))
        print(name, "| paragraphs:", len(d.paragraphs), "| first:", d.paragraphs[0].text)
