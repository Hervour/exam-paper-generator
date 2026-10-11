# -*- coding: utf-8 -*-
"""Convert the exam-paper markdown files to Word (.docx) with Chinese exam formatting."""
import os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"C:\Users\LHW66\Documents\Codex\2026-08-19\wo\work\md"
DST = r"C:\Users\LHW66\Documents\Codex\2026-08-19\wo\work\docx"
os.makedirs(DST, exist_ok=True)

EA_BODY = "宋体"
EA_HEAD = "黑体"
LATIN = "Times New Roman"

def set_run_font(run, size, bold=False, ea=EA_BODY, color=None):
    run.font.name = LATIN
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), LATIN)
    rFonts.set(qn("w:hAnsi"), LATIN)
    rFonts.set(qn("w:eastAsia"), ea)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def set_first_line_indent_chars(paragraph, chars=200):
    pPr = paragraph._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars))
    ind.set(qn("w:firstLine"), "420")

def add_rich_text(paragraph, text, size=10.5, base_bold=False, ea=EA_BODY, color=None):
    """Parse **bold** markers and add runs."""
    tokens = re.split(r"(\*\*.+?\*\*)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = paragraph.add_run(tok[2:-2])
            set_run_font(run, size, bold=True, ea=ea, color=color)
        else:
            run = paragraph.add_run(tok)
            set_run_font(run, size, bold=base_bold, ea=ea, color=color)

def make_paragraph(doc, text, size=10.5, bold=False, ea=EA_BODY, align=None,
                   indent=True, space_after=3, color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.3
    if align is not None:
        p.alignment = align
    if indent:
        set_first_line_indent_chars(p)
    add_rich_text(p, text, size=size, base_bold=bold, ea=ea, color=color)
    return p

def add_table(doc, rows):
    """rows: list of list of cell text."""
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            text = row[j] if j < len(row) else ""
            run = p.add_run(text)
            set_run_font(run, 10.5, bold=(i == 0))
    # spacing after table
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def convert(md_path, docx_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    # A4 page
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped == "---":
            doc.add_page_break()
            i += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):  # skip separator row
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        if stripped.startswith("# "):
            text = stripped[2:].strip()
            p = make_paragraph(doc, text, size=16, bold=True, ea=EA_HEAD,
                               align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=2)
            i += 1
            continue
        if stripped.startswith("## "):
            text = stripped[3:].strip()
            p = make_paragraph(doc, text, size=14, bold=True, ea=EA_HEAD,
                               align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, space_after=6)
            p.paragraph_format.space_before = Pt(10)
            i += 1
            continue
        if stripped.startswith("### "):
            text = stripped[4:].strip()
            p = make_paragraph(doc, text, size=12, bold=True, ea=EA_HEAD,
                               align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, space_after=3)
            i += 1
            continue
        if stripped.startswith("> "):
            text = stripped[2:].strip()
            if text.startswith("考试时间"):
                p = make_paragraph(doc, text, size=12, align=WD_ALIGN_PARAGRAPH.CENTER,
                                   indent=False, space_after=2)
            else:
                p = make_paragraph(doc, text, size=10, indent=False, space_after=6, color=(0x60, 0x60, 0x60))
            i += 1
            continue
        # emphasis-only line like **（难度偏高 ...）**
        if stripped.startswith("**") and stripped.endswith("**") and stripped.count("**") == 2:
            p = make_paragraph(doc, stripped, size=12, align=WD_ALIGN_PARAGRAPH.CENTER,
                               indent=False, space_after=6)
            i += 1
            continue
        # body paragraph
        make_paragraph(doc, stripped, size=10.5)
        i += 1

    doc.save(docx_path)
    print("saved:", os.path.basename(docx_path))

files = {
    "数学 分班考试模拟卷（难度偏高）.md": "数学 分班考试模拟卷（难度偏高）.docx",
    "物理 分班考试模拟卷（难度偏高）.md": "物理 分班考试模拟卷（难度偏高）.docx",
    "化学 分班考试模拟卷（难度偏高）.md": "化学 分班考试模拟卷（难度偏高）.docx",
}
for src, dst in files.items():
    convert(os.path.join(SRC, src), os.path.join(DST, dst))
print("ALL DONE")

